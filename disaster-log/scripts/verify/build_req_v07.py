# -*- coding: utf-8 -*-
"""재난대응·복구 상세 기능 요구사항 정의서 v0.7 — 개발 반영본. 서식은 plan/exel style.xlsx 기준."""
import sys, json
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.utils import get_column_letter

sys.stdout.reconfigure(encoding="utf-8")
from pathlib import Path
ROOT = Path(__file__).resolve().parents[3]  # 저장소 루트 (une_report3)
OUT = str(ROOT / "plan" / "재난대응_복구_상세기능요구사항정의서_v0.7.xlsx")
SRC = str(ROOT / "plan" / "_이전버전" / "req_v06.json")  # 원본 v0.6 행 JSON
DATE = "2026-09-14"

# ── 서식 (exel style.xlsx) ──
FONT = "Pretendard"
INK = "FF1A1A1A"; GRAY = "FF6B6B6B"; WHITE = "FFFFFFFF"
F_BODY = Font(name=FONT, size=10, color=INK)
F_BOLD = Font(name=FONT, size=10, bold=True, color=INK)
F_GRAY = Font(name=FONT, size=10, color=GRAY)
F_HEAD = Font(name=FONT, size=10, bold=True, color=WHITE)
F_TITLE = Font(name=FONT, size=20, bold=True, color=INK)
F_SUB = Font(name=FONT, size=10, color=GRAY)
FILL_HEAD = PatternFill("solid", fgColor="FF1A1A1A")
FILL_LABEL = PatternFill("solid", fgColor="FFF7F7F7")
FILL_BEIGE = PatternFill("solid", fgColor="FFF4F1EA")
FILL_NEW = PatternFill("solid", fgColor="FFEAF7EE")    # 연두 = v0.7 신규
FILL_MOD = PatternFill("solid", fgColor="FFFFF9E8")    # 연노랑 = v0.7 수정·보완
thin = Side(style="thin", color="FFD9D9D9")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
BORDER_B = Border(bottom=thin)
AL_TOP = Alignment(horizontal="left", vertical="top", wrap_text=True)
AL_TOP_NOWRAP = Alignment(horizontal="left", vertical="top")
AL_CENTER = Alignment(horizontal="center", vertical="top")
AL_RIGHT = Alignment(horizontal="right", vertical="top")
AL_HEAD = Alignment(horizontal="left", vertical="center")

def header_row(ws, row, headers, widths):
    for i, (h, w) in enumerate(zip(headers, widths), start=1):
        c = ws.cell(row, i, h); c.font = F_HEAD; c.fill = FILL_HEAD; c.border = BORDER_B; c.alignment = AL_HEAD
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[row].height = 26

def est_height(texts, widths):
    """열 너비 기준 줄 수 추정 → 행 높이"""
    lines = 1
    for t, w in zip(texts, widths):
        if not t: continue
        s = str(t); n = 0
        for part in s.split("\n"):
            n += max(1, int(len(part) * 1.9 / max(w, 6)) + 1)
        lines = max(lines, n)
    return min(400, max(32, 14 * lines + 6))

# ── 요구사항 데이터 ──
# 상태: 구현 / 부분 / 모의 / 계획
# 각 항목: (영역ID, 영역, 요구ID, 요구명, 내용, 입력, 출력, 연계, 상태, 구현 화면·모듈, 비고, 변경표시 None|new|mod)
old = json.load(open(SRC, encoding="utf-8"))
OLD = {r[2]: r for r in old}

def base(rid, status, module, note_add=None, change=None, content=None, name=None, inp=None, out=None, link=None):
    r = OLD[rid]
    return (r[0], r[1], rid, name or r[3], content or r[4], inp or r[5], out or r[6], link or r[7], status, module,
            (note_add + " | " if note_add else "") + (r[8] or ""), change)

def new(area_id, area, rid, name, content, inp, out, link, status, module, note):
    return (area_id, area, rid, name, content, inp, out, link, status, module, note, "new")

ROWS = [
  # ── UFR-001 상황관리 ──
  base("UFR-001-001", "구현", "새 업무 시작(/situations/new) · useAppStore.createSituation"),
  base("UFR-001-002", "구현", "새 업무 시작 · 지자체·재난유형(풍수해 호우/태풍/대설·산불) 선택"),
  base("UFR-001-003", "구현", "새 업무 시작 · 시군구·읍면동 복수 지역 추가/삭제"),
  base("UFR-001-004", "구현", "새 업무 시작 · 현재상황 자유입력 → AI Context(buildContext)"),
  base("UFR-001-005", "구현", "대시보드 업무 목록 · localStorage persist · 설정 › JSON 전체 내보내기/불러오기",
       note_add="v0.7: DB 없이 브라우저 localStorage(zustand persist)에 저장. 설정 화면 JSON 백업/복원 제공", change="mod"),
  base("UFR-001-006", "구현", "새 업무 시작(훈련) · 훈련명·목적·시나리오·참여기관·일정"),
  base("UFR-001-007", "구현", "상황·기상·상황부여 탭 · addInjection → 원장 injection 이벤트"),
  new("UFR-001", "상황관리", "UFR-001-008", "대시보드·진행 중 업무 현황",
      "시스템은 진행 중 업무(실제재난/훈련) 목록, 완료 조치·SMS·자원·누적 이벤트 요약, SOP 라이브러리 현황, 업무흐름 8단계 안내를 대시보드로 제공하고 업무 카드에서 바로 이어서 진행할 수 있어야 한다.",
      "저장된 업무·SOP 라이브러리", "대시보드 요약·업무 카드·최근 업무(LNB)", "UNE", "구현", "대시보드(/) · AppShell LNB 「진행 중 업무」",
      "v0.7 신규. 부산 풍수해 Seed·안전한국훈련 Seed 를 한 번에 불러오는 버튼 포함"),
  # ── UFR-002 기상정보 ──
  base("UFR-002-001", "모의", "상황·기상 탭 · seed/weather.mockWeatherAlerts",
       note_add="v0.7: 기상청 공식 API 미연계. 지자체·재난유형별 Seed 특보를 동일 인터페이스로 표시", change="mod"),
  base("UFR-002-002", "모의", "lib/t3q/adapter.weatherSummary (Mock)",
       note_add="v0.7: T3Q 실 API 미사용(사용자 지시). 동일 시그니처의 Mock 어댑터로 요약문 생성, 「확인 필요」 표시", change="mod"),
  base("UFR-002-003", "구현", "WeatherAlert.source/issuedAt · WeatherSummary(출처·조회시각·verify)"),
  # ── UFR-003 문서 ──
  base("UFR-003-001", "모의", "문서·조치 선택 탭 · mockT3Q.searchDocuments → 파일명 목록",
       note_add="v0.7: T3Q Mock. 부산 풍수해·산불·훈련 가이드 등 Seed 문서 메타(documents.ts)에서 Context 매칭", change="mod"),
  base("UFR-003-002", "구현", "문서·조치 선택 탭 · toggleDoc (파일명만 표시)"),
  base("UFR-003-003", "구현", "실행 탭 「관련 문서 재조회 · 추가 SOP」 → 문서·조치 선택 탭 복귀"),
  # ── UFR-004 SOP ──
  base("UFR-004-001", "모의", "mockT3Q.recommendActions · 문서 선택 후 1회 호출 → 조치목록+상세",
       note_add="v0.7: Mock 어댑터. 풍수해 113개·산불 10개 조치 카드(flood_actions.json, wildfire.ts)", change="mod"),
  base("UFR-004-002", "구현", "RecommendedAction 타입(코드·조치명·단계·주관/지원/협업·세부행동·전파대상·필요자원·출처문서·페이지)"),
  base("UFR-004-003", "구현", "문서·조치 선택 탭 조치 카드(단계·담당 요약, 펼치면 세부행동)"),
  base("UFR-004-004", "구현", "toggleAction · 체크 선택/해제"),
  base("UFR-004-005", "구현", "moveAction(위/아래) · 선택 순서 목록"),
  base("UFR-004-006", "구현", "lib/sop/adapter.actionsToCompns → compnsToFlow · CompnSaveParams(CompnTyCode 104001~104006) 계약 유지"),
  base("UFR-004-007", "구현", "toNode: leadDept/supportDept/coopAgencies/details/targets/resources → SopNodeData"),
  base("UFR-004-008", "구현", "SOP 구성·편집 탭 · ReactFlow 캔버스(@xyflow) · NodePanel 분기값·연결대상 편집"),
  base("UFR-004-009", "구현", "SopVersion(recommended/edited/confirmed) · confirmSop · 버전 이력 패널 · setActiveSop"),
  new("UFR-004", "SOP 추천·생성·편집", "UFR-004-010", "SOP 라이브러리(상황 독립 작성·게시·버전)",
      "사용자는 상황이 없어도 SOP 를 미리 작성·편집(초안)하고 「게시」하여 실행 가능 버전으로 확정할 수 있어야 한다. 시스템은 재난유형·태그·설명·게시 이력·배포 이력을 관리하고 목록에서 검색·복제·삭제·JSON 내보내기/가져오기를 제공해야 한다.",
      "SOP 이름·재난유형·태그·노드/연결·게시 메모", "SopTemplate(초안/게시본 v1..n/이력/사용)", "UNE", "구현",
      "SOP 관리(/sops) · SOP 편집기(/sops/[id]) · createTemplate/publishTemplate/duplicateTemplate/importTemplate",
      "v0.7 신규. 편집(초안)과 실행(배포 스냅샷)을 분리. 기본 3종(호우 초기대응·태풍 주민 사전대피·산불 초동조치) Seed"),
  new("UFR-004", "SOP 추천·생성·편집", "UFR-004-011", "라이브러리 SOP 상황 배포·즉시 실행",
      "상황 발생 시 사용자는 라이브러리 게시본을 선택해 해당 상황의 실행본으로 배포하고 필요하면 즉시 실행을 시작할 수 있어야 한다. 배포는 그 시점의 스냅샷이며 이후 라이브러리 수정이 진행 중 상황에 영향을 주지 않아야 한다. 상황에서 편집한 SOP 를 라이브러리 초안으로 되돌려 저장할 수 있어야 한다.",
      "상황 ID, 라이브러리 SOP, 즉시실행 여부", "SopVersion(kind=confirmed, templateId/templateVersion) · 라이브러리 usage", "UNE", "구현",
      "SOP 구성·편집 탭 「라이브러리에서 배포」/「라이브러리에 저장·반영」 · 실행 탭 빈 상태 「라이브러리에서 선택해 바로 실행」 · deployTemplate/pushToLibrary",
      "v0.7 신규. 재난유형 일치 게시본 우선 표시(LibraryPicker)"),
  new("UFR-004", "SOP 추천·생성·편집", "UFR-004-012", "AI SOP 자유생성(UNI RAG)",
      "사용자는 자연어로 상황을 설명해 UNI RAG 시스템(/chat/json)에 SOP 컴포넌트 생성을 요청할 수 있어야 하며, 시스템은 SSE 스트림(__status__/__compn__/__done__)을 받아 기존 SOP 컴포넌트 계약(CompnSaveParams)으로 캔버스에 실시간 반영해야 한다. 실행 중인 상황에서는 활성 실행본을 덮어쓰지 않아야 한다.",
      "자연어 질의, 재난유형·지자체 Context", "생성 SOP 버전(추천 원본) · 생성 상태·모델·출처(uni/fallback)", "UNI RAG · UNE", "부분",
      "SopTab/SOP 편집기 「AI 생성」 모달 · /api/uni/sop SSE 릴레이 · lib/uni/client(로그인 토큰 캐시·모델 선택·circuit breaker)",
      "v0.7 신규. UNI 서버(사내망)는 개발 PC·Vercel 에서 미도달 → 실호출 미검증, 접속 불가 시 Seed 기반 대체 생성. SignalR 릴레이는 SSE 직접 릴레이로 치환"),
  new("UFR-004", "SOP 추천·생성·편집", "UFR-004-013", "SOP 노드 상황전파 채널 설정",
      "사용자는 SOP 노드 속성에서 상황전파 채널(SMS·이메일)을 지정할 수 있어야 하며, 실행 중 상황전파 발송의 기본 채널로 사용되어야 한다.",
      "노드, 채널(SMS/이메일)", "SopNodeData.channels", "UNE", "구현", "NodePanel 「상황전파 채널」 체크박스", "v0.7 신규. 미지정 시 SMS"),
  new("UFR-004", "SOP 추천·생성·편집", "UFR-004-014", "SOP 편집 속성 패널 숨기기·폭 조절",
      "사용자는 SOP 편집 화면의 우측 속성 패널을 숨기거나 다시 열 수 있어야 하고, 패널 가장자리를 드래그해 폭(280~720px)을 조절할 수 있어야 한다. 설정은 브라우저에 기억되어야 한다.",
      "드래그·토글", "패널 표시 상태·폭", "UNE", "구현", "components/sop/SidePanel (SOP 구성·편집 탭 · 라이브러리 편집기 공용)", "v0.7 신규. 더블클릭 시 기본 폭"),
  # ── UFR-005 실행 ──
  base("UFR-005-001", "구현", "실행 탭 · startRun/stopRun · 실행 시작 확인 → 「실행이 시작되었습니다」 안내 · 종료 노드 도달 시 완료 안내 · 중지 확인 모달",
       note_add="v0.7: 시작·종료·중지 시 사용자가 확실히 인지할 수 있는 안내창 추가. 다른 화면에서 시작(라이브러리 즉시실행)해도 탭 진입 시 안내", change="mod"),
  base("UFR-005-002", "구현", "실행 탭 좌 기준 SOP(목록/플로우 보기) · 우 조치 카드 + 실제 수행 타임라인",
       note_add="v0.7: 플로우 보기에 실행 상태·현재 조치 표시, 노드 클릭으로 처리", change="mod"),
  base("UFR-005-003", "구현", "NodeRun(startedAt/finishedAt/status) · startNode/completeNode/skipNode · 다음 노드 자동 진행"),
  base("UFR-005-004", "구현", "조치 카드 「현장메모」 → 원장 memo 이벤트"),
  base("UFR-005-005", "구현", "조치 카드 「조치결과」 → 원장 result 이벤트 · 완료 후 결과 수정 저장"),
  base("UFR-005-006", "구현", "조치 카드 「실제 담당자」(사용자 기본값·수정 가능)"),
  base("UFR-005-007", "구현", "조치 카드 「증빙 첨부」(파일명·크기 기록)", note_add="파일 본문은 저장하지 않고 메타만 기록(DB 부재)", change="mod"),
  base("UFR-005-008", "구현", "실행 탭 「상황 변화 대응」 카드 · 관련 문서 재조회 → 추가 SOP 버전 · 상황변화 메모 모달",
       note_add="v0.7: 「상황 변화 대응」 용어 도움말(?) 추가, 변화 메모는 원장 user 이벤트", change="mod"),
  base("UFR-005-009", "구현", "훈련 모드 임무 수신/확인/완료 버튼(missionAck) + 모바일 상황전파 응답(UFR-006-011) 연동",
       note_add="v0.7: 회신 방식을 「인증 없는 모바일 웹 링크」로 확정·구현. 실제재난에서도 동일 이력 사용", change="mod"),
  base("UFR-005-010", "구현", "상황판단 노드 분기값 ChoiceChip → completeNode(branchValue) → nextNodes(label 일치)"),
  new("UFR-005", "SOP 실행·조치관리", "UFR-005-011", "세부행동 체크 수행기록",
      "조치 카드의 세부행동(매뉴얼 자동입력)은 체크박스로 제공되어야 하며, 사용자가 수행한 항목을 체크하면 체크 시각과 함께 원장에 「세부행동 수행」 기록이 남고 목록에 진행률(n/m)이 표시되어야 한다.",
      "세부행동 체크/해제", "NodeRun.checks(index→시각) · 원장 run 이벤트", "UNE", "구현", "실행 탭 조치 카드 「세부행동 체크」 · toggleDetailCheck", "v0.7 신규. 체크 해제 시 해당 기록 제거"),
  new("UFR-005", "SOP 실행·조치관리", "UFR-005-012", "조치결과 자동 정리",
      "조치결과를 비워 두고 완료하면 체크한 세부행동을 조치결과 문장으로 자동 정리해야 하며, 체크도 결과도 없으면 확인 모달을 거쳐 결과 없이 완료할 수 있어야 한다.",
      "체크 항목, 조치결과", "정리된 조치결과", "UNE", "구현", "조치 카드 「체크한 세부행동 n건을 조치결과로 채우기」 · complete()", "v0.7 신규"),
  new("UFR-005", "SOP 실행·조치관리", "UFR-005-013", "화면 용어 도움말",
      "지자체 담당자에게 낯선 용어(기준 SOP·실제 수행 타임라인·상황 변화 대응·세부행동·분기값·실행본·SOP 준비 방법·조직 구성 등)에 ? 아이콘을 두고 마우스 오버 시 설명을 제공해야 한다.",
      "마우스 오버", "설명 툴팁", "UNE", "구현", "components/ui Help(DS Tooltip) · 실행·SOP·조직·상황전파 화면", "v0.7 신규"),
  # ── UFR-006 상황전파·재난자원 ──
  base("UFR-006-001", "구현", "상황전파 발송 모달 · 채널·제목·내용 편집 · 수신자(조직도 사람/전송그룹/조치 담당 추천/검색)",
       name="상황전파 채널·수신자·문안 편집",
       content="사용자는 SOP 조치에서 상황전파 채널(SMS·이메일), 제목·내용, 수신자를 확인·수정할 수 있어야 한다. 수신자는 조직도(지휘부·실무반·유관기관 → 부서 → 사람)에서 사람 단위로, 주소록 전송그룹 단위로 선택할 수 있어야 하며, 조치의 주관·지원·협업 부서 인원은 추천으로 제시되어야 한다.",
       inp="채널, 제목, 내용, 사람/그룹 선택", out="발송 대상·문안·채널",
       note_add="v0.7: SMS 문안 편집 → 채널·수신자(조직/그룹) 선택으로 확장. T3Q 전파대상은 참고 표시", change="mod"),
  base("UFR-006-002", "모의", "addDispatch → 수신자별 토큰 링크·QR·SMS 문안 생성 · 실제 발송 자리(UNE SMS 모듈·메일 서버)",
       name="상황전파 발송(SMS·이메일)",
       content="사용자는 선택한 채널로 상황전파를 발송할 수 있어야 한다. 발송 메시지에는 현장 응답용 모바일 링크가 포함되어야 하며, 실제 문자·메일 발송은 UNE 기존 SMS 모듈·메일 서버 연계로 수행한다.",
       inp="mode, 채널, 수신자, 제목·내용", out="Dispatch(수신자별 토큰·링크), SmsRecord, 발송 결과",
       note_add="v0.7: 사용자 전제(SMS·이메일을 보낸다고 가정)에 따라 모의 발송. 링크·QR·SMS 문안 생성까지 구현, 실발송 연계는 계획", change="mod"),
  base("UFR-006-003", "구현", "Situation.dispatches(수신자별 수신/완료/조치사항) + sms(호환) · 원장 sms 이벤트 · 결과보고 붙임 1·1-1",
       note_add="v0.7: 발송이력에 채널·수신자별 응답 상태 포함", change="mod"),
  base("UFR-006-004", "구현", "자원 투입 모달(수기/내부 자원목록 · 재난유형별 목록) · returnResource(회수)"),
  base("UFR-006-005", "구현", "자원 투입 모달 KRMS 탭(비활성 · 연계 자리 안내)"),
  base("UFR-006-006", "구현", "ResourceRecord.source(manual/internal/KRMS)"),
  new("UFR-006", "상황전파·재난자원", "UFR-006-007", "조직·연락처 관리(재대본 편성 3계층)",
      "시스템은 상황전파 수신대상 연락망을 지자체 재난안전대책본부 편성에 맞춰 관리해야 한다. 구조는 구분(지휘부 · 13개 협업기능별 실무반 · 유관기관) → 실무반·기관 → 부서 → 사람(직위·이름·전화번호·이메일·비고)이며, 조직도·목록 보기, 개별 등록·수정·삭제, 검색·구분 필터를 제공해야 한다.",
      "구분·실무반/기관·부서명·직위·이름·전화번호·이메일·비고", "Contact 목록 · 조직도", "UNE", "구현",
      "설정 › 조직·연락처 관리(/settings/org) 「조직도」「연락처 목록」 · contacts 스토어",
      "v0.7 신규. 부산 풍수해 매뉴얼 p.19 지휘부 · 붙임 8-5-① 13개 실무반·담당부서 · p.23 유관기관 기준. 조치의 주관/지원/협업은 조치별 역할이므로 SOP 노드에 기록"),
  new("UFR-006", "상황전파·재난자원", "UFR-006-008", "연락처 엑셀 양식 다운로드·일괄 업로드",
      "사용자는 연락처 업로드용 엑셀 양식(연락처 시트 + 작성안내 시트)을 내려받아 작성 후 일괄 업로드할 수 있어야 한다. 시스템은 헤더 별칭(부서/소속·직급·성명·연락처 등)을 자동 인식하고, 구분·실무반은 병합셀처럼 첫 행 값을 이어받으며, 미리보기에서 오류·경고를 표시한 뒤 추가/병합/전체교체를 선택해 반영해야 한다. 현재 목록을 동일 양식으로 내보낼 수 있어야 한다.",
      ".xlsx/.xls/.csv 파일", "검증 결과 미리보기 · 반영된 연락처", "UNE(SheetJS)", "구현",
      "조직·연락처 관리 「엑셀 양식 다운로드」「엑셀 일괄 업로드」「내보내기」 · lib/org/excel.ts", "v0.7 신규. 전화번호 자동 정규화(000-0000-0000)"),
  new("UFR-006", "상황전파·재난자원", "UFR-006-009", "전송그룹(주소록) 관리",
      "사용자는 자주 함께 전파하는 수신자 묶음을 전송그룹으로 등록·편집·삭제할 수 있어야 하며, 구성원은 조직도에서 부서·실무반 단위 체크 또는 검색으로 담을 수 있어야 한다. 상황전파 발송 시 그룹 단위로 선택할 수 있어야 한다.",
      "그룹명·설명·구성원", "SendGroup 목록", "UNE", "구현", "조직·연락처 관리 「전송그룹」 탭 · groups 스토어 · ContactPicker",
      "v0.7 신규. 기본 5개(재대본 지휘부·13개 협업기능반 반장·재난상황관리반 전원·호우 지원부서·유관기관 상황실) Seed"),
  new("UFR-006", "상황전파·재난자원", "UFR-006-010", "현장 요원 모바일 상황전파 페이지",
      "상황전파 수신자는 인증 없이 링크로 열리는 모바일 웹페이지에서 상황·조치·전파 내용·수신자 정보를 확인하고 ① 수신확인 → ② 임무완료 → ③ 조치사항(선택) 순으로 응답할 수 있어야 한다. 페이지는 필요한 정보만 담고 반응형(갤럭시+크롬 기준, iOS 사파리 표준 CSS)이며 큰 버튼으로 구성되어야 한다.",
      "토큰 링크(payload 포함), 응답 조작", "응답(수신확인·임무완료·조치사항·시각)", "UNE", "구현",
      "/m/[token] · components/dispatch/MobileDispatchView · payload base64url(링크만으로 렌더) · /api/dispatch/[token]",
      "v0.7 신규. 응답 상태는 기기 localStorage 에도 보관해 새로고침 후 유지. iOS 실기기 검증은 미실시"),
  new("UFR-006", "상황전파·재난자원", "UFR-006-011", "현장 응답 수신 및 실행내역·상황일지 반영",
      "시스템은 현장 요원의 응답을 수신하여 해당 조치의 응답 현황(수신 n/총·완료 n/총·수신자별 시각·조치사항), 실제 수행 타임라인, 상황 이벤트 원장(mission/result), 노드 실행기록(missionAck), 결과보고 붙임표에 자동 반영해야 한다. 조치사항은 「확인 필요」 상태로 들어와 담당자가 확정한다.",
      "응답 이벤트(토큰·종류·시각·조치사항)", "Dispatch 수신자 상태 · 원장 이벤트 · missionAck", "UNE 서버 저장소(메모리/Upstash Redis)", "구현",
      "lib/dispatch/useDispatchSync(4초 폴링 + BroadcastChannel) · applyDispatchAck · DispatchStatus · 결과보고 붙임 1-1",
      "v0.7 신규. 운영(Vercel)은 Upstash Redis(마켓플레이스 연동) 저장, 로컬도 동일 DB 사용. 7일 보관"),
  new("UFR-006", "상황전파·재난자원", "UFR-006-012", "상황전파 모바일 미리보기·링크·QR",
      "상황전파 발송 화면은 현장 요원이 볼 모바일 페이지를 발송 전에 미리 보여 주어야 하며, 발송 후에는 수신자별 링크·QR·SMS 문안을 제공해 복사하거나 「모의 수신」으로 열어 볼 수 있어야 한다.",
      "발송 초안", "휴대폰 프레임 미리보기 · 수신자별 링크·QR·문안", "UNE(qrcode)", "구현", "DispatchModal PhoneFrame · QRCode.toDataURL", "v0.7 신규. 실발송 시 URL 단축 서비스 병행 권장"),
  # ── UFR-007 상황일지 ──
  base("UFR-007-001", "구현", "LedgerEvent 원장(situation/weather/document/sop/run/result/memo/sms/resource/injection/mission/branch/log/report/user) · 모든 액션에서 자동 누적",
       note_add="v0.7: 세부행동 수행·상황전파 발송·현장 응답(수신확인/임무완료/조치사항) 이벤트 추가", change="mod"),
  base("UFR-007-002", "구현", "상황일지 탭 좌측 원장 타임라인 · 유형 필터 칩 · 확인/확인 필요/제외 토글"),
  base("UFR-007-003", "부분", "상황일지 탭 「AI 초안 생성」 · /api/uni/chat(UNI 일반 챗 SSE) · 접속 불가 시 lib/uni/fallback 규칙 기반 초안(출처 fallback 표시)",
       name="상황일지 AI 초안 생성(UNI)",
       content="사용자는 확인된 이벤트를 바탕으로 UNI RAG 일반 챗 API 에 실제재난 또는 훈련 상황일지 초안 생성을 요청할 수 있어야 한다. 초안은 Markdown(일자별 ### · 시각별 불릿 · 종합) 형식이어야 하며, 서버 접속 불가 시에도 동일 형식의 대체 초안을 제공해야 한다.",
       link="UNI RAG(일반 챗) · 대체: UNE 규칙 생성",
       note_add="v0.7: 사용자 지시로 T3Q → UNI 일반 챗 API 로 변경. UNI 실호출 미검증(사내망), KST 시각 처리", change="mod"),
  base("UFR-007-004", "구현", "상황일지 탭 편집/분할/미리보기(react-markdown) · 확정 시 이력 저장",
       note_add="v0.7: 본문은 Markdown 편집·미리보기", change="mod"),
  base("UFR-007-005", "구현", "LOG_TEMPLATES(실제재난 표준/훈련) 선택 → 초안 Context"),
  # ── UFR-008 결과보고 ──
  base("UFR-008-001", "구현", "결과보고 탭 · 보고 구분(mode 기준 TOC)·제목·기준시각"),
  base("UFR-008-002", "구현", "REPORT_TOC 공통 목차 · 섹션 추가/삭제/순서/제목 편집·포함 여부"),
  base("UFR-008-003", "구현", "결과보고 「반영자료」 체크(results/log/sms/resources/weather/damage/photos/plan/injections/missions)"),
  base("UFR-008-004", "구현", "섹션 본문 직접 입력(Markdown)"),
  base("UFR-008-005", "구현", "lib/export/build.reportExportDoc · 조치결과·SMS·상황전파 응답·자원 현황표 자동 생성"),
  base("UFR-008-006", "부분", "결과보고 「AI 본문 초안 생성(UNI 챗)」 · 섹션별 Markdown 초안 · 대체 생성",
       name="문안 정리·본문 초안 연계(UNI)",
       content="사용자는 필요 시 목차별 본문 초안을 UNI 일반 챗으로 생성·정리할 수 있어야 하며, 미연계 상태에서도 규칙 기반 대체 초안으로 보고서 작성·출력이 가능해야 한다.",
       link="UNI RAG(일반 챗) · 대체: UNE 규칙 생성",
       note_add="v0.7: T3Q 협의 대신 UNI 챗 적용. 실호출 미검증", change="mod"),
  base("UFR-008-007", "구현", "결과보고 탭 섹션별 Markdown 편집/미리보기 · 전체 미리보기"),
  base("UFR-008-008", "구현", "HWPX(templete/AI 행정문서 템플릿.hwpx 스타일 ID 매핑 자체 생성기) · DOCX(docx) · PDF(인쇄)",
       name="문서 내보내기(HWPX 행정문서 양식·DOCX·PDF)",
       content="사용자는 최종 상황일지 및 결과보고서를 한글(HWPX)·Word(DOCX)·PDF 로 내보낼 수 있어야 한다. HWPX 는 지정된 행정문서 템플릿의 문단·글자 스타일(제목·부제·장제목·본문·캡션·표)을 적용해야 한다.",
       out="HWPX/DOCX/PDF",
       note_add="v0.7: 지원형식 확정. HWPX 는 XML 정합성 검증 완료, 한컴오피스 실제 열기 확인 필요", change="mod"),
  # ── UFR-009 공통 ──
  base("UFR-009-001", "구현", "SourceKind(user/official/ai/system/sms/resource/sop) · VerifyState(confirmed/unverified/excluded) · 배지 표시"),
  base("UFR-009-002", "구현", "SopVersion 이력 · log/report history · 원장 sop/log/report 이벤트"),
  base("UFR-009-003", "구현", "UNI 접속 실패 시 circuit breaker(60초) + 대체 생성 · 헤더 「UNI 대체모드」 배지 · 설정 화면 연계 상태",
       note_add="v0.7: UNI/T3Q 장애 시 대체 경로 구현. 핵심 업무(SOP 구성·실행·기록·출력)는 외부 연계 없이 동작", change="mod"),
  base("UFR-009-004", "구현", "Seed: 부산 풍수해(113 조치) · 안전한국훈련(태풍) · 경북 산불(10 조치) · 재난유형별 문서·기상·자원목록",
       note_add="v0.7: 3개 검증축 모두 데모 Seed 및 시연 시나리오로 구현", change="mod"),
  new("UFR-009", "공통", "UFR-009-005", "시연모드(자동 시연)",
      "시스템은 헤더의 「시연모드」로 산불 실제재난 시나리오 15단계(상황 등록 → 문서·조치 → SOP → 라이브러리 → 실행·세부행동·상황전파·현장 응답·상황 변화 → 완료 → 상황일지 → 결과보고)를 단계별로 자동 수행하며, 모달이 아닌 화면을 가리지 않는 플로팅 패널(접기·좌/우 도킹·자동 진행·단축키·단계 목록)로 화면을 제어하고 발표용 내레이션·포인트·「직접 해보기」를 보여 주어야 한다.",
      "다음/이전/단계 선택", "화면 이동 + 데이터 조작 + 내레이션", "UNE", "구현",
      "AppShell 「시연모드」 · components/demo/DemoPanel · lib/demo/scenario · useDemoStore(sessionStorage)",
      "v0.7 신규. 마지막 단계에서 시연 상황 삭제 후 종료 가능"),
  new("UFR-009", "공통", "UFR-009-006", "사내 디자인시스템 적용",
      "화면은 ㈜유엔이 디자인시스템(@une-front/react-ui · design-tokens)의 컴포넌트·토큰·아이콘을 사용해 구성해야 하며, 디자인시스템에 없는 요소(SOP 노드 아이콘 등)만 대체 아이콘을 사용한다.",
      "-", "일관된 UI", "UNE DS", "구현", "components/ui(어댑터) · components/icons · globals.css(1rem=1px 스케일)", "v0.7 신규. DS 배포판 미포함 아이콘(IconNode*)은 lucide 대체 → 차기 DS 버전에서 교체"),
  new("UFR-009", "공통", "UFR-009-007", "저장·배포 구조(DB 없음)",
      "시스템은 서버 DB 없이 브라우저 저장소(localStorage)와 서버 라우트(외부 API 키 보호·SSE 릴레이·현장 응답 저장소)만으로 동작하고 Vercel 에 배포 가능해야 한다. 상황 데이터는 JSON 으로 내보내기/불러오기 할 수 있어야 한다.",
      "-", "배포 사이트 · JSON 백업", "Vercel · Upstash Redis", "구현",
      "https://une-report3.vercel.app · vercel.json · 설정 › 데이터 백업", "v0.7 신규. 여러 PC 가 한 상황을 공유하려면 향후 서버 DB(Postgres 등) 도입 필요"),
  new("UFR-009", "공통", "UFR-009-008", "UNI RAG 연계 상태 표시",
      "시스템은 UNI RAG 서버 연결 여부와 사용 모델을 주기적으로 확인해 헤더 배지와 설정 화면에 표시해야 하며, 인증정보는 서버 환경변수에서만 처리하고 브라우저에 노출하지 않아야 한다.",
      "환경변수(UNI_BASE_URL/ACCOUNT/PASSWORD 등)", "연결 상태·모델·오류", "UNI RAG", "구현", "/api/uni/health · AppShell UNI 배지 · 설정·연계상태 화면", "v0.7 신규. 사내망 외부(Vercel)에서는 항상 대체모드"),
]

# ── 통계 ──
n_new = sum(1 for r in ROWS if r[11] == "new"); n_mod = sum(1 for r in ROWS if r[11] == "mod")
st = {}
for r in ROWS: st[r[8]] = st.get(r[8], 0) + 1
print("total", len(ROWS), "new", n_new, "mod", n_mod, st)

wb = Workbook()
# ── 시트 1: 작성안내 ──
ws = wb.active; ws.title = "작성안내"
ws.column_dimensions["A"].width = 14; ws.column_dimensions["B"].width = 120
ws.cell(1, 1, "재난대응·복구 상세 기능 요구사항 정의서 v0.7 (개발 반영본)").font = F_TITLE
ws.merge_cells("A1:B1"); ws.row_dimensions[1].height = 34
for c in (1, 2): ws.cell(1, c).border = BORDER_B
ws.cell(2, 1, f"㈜유엔이 · 재난상황일지 생성도구(disaster-log) 개발 결과 반영 · v0.7 ({DATE}) · 기준 v0.6 (2026-09-13)").font = F_SUB
ws.merge_cells("A2:B2"); ws.row_dimensions[2].height = 20
guide = [
    ("목적", "v0.6 요구사항 55건을 실제 개발 결과(프로토타입)와 대조하여 구현 상태·구현 화면·모듈을 기록하고, 개발 과정에서 추가된 기능(SOP 라이브러리, 조직·연락처, 전송그룹, 상황전파 모바일 응답, 세부행동 체크, 시연모드 등)을 신규 요구사항으로 정의한다.", 48),
    ("기준", "실제재난·안전한국훈련 공통 Core / 문서 선택 → 조치 상세 일괄 수신 → UNE SOP 구성 → 실행·조치결과 → 상황일지 → 결과보고. T3Q 실 API 는 미사용(Mock 어댑터, 동일 인터페이스), AI 는 사내 UNI RAG(SOP: /chat/json, 초안: 일반 챗) 사용, 기존 UNE SOP 자동생성 모듈 컴포넌트 계약(CompnSaveParams) 재사용.", 64),
    ("구현 기준", "Next.js 16 · React 19 · 사내 디자인시스템 @une-front · zustand(localStorage, DB 없음) · Vercel 배포(https://une-report3.vercel.app) · GitHub jazzsalle/une_report3. 조직 구성은 부산광역시 풍수해 재난 현장조치 행동매뉴얼(26.7.) 재난안전대책본부 편성(지휘부 · 13개 협업기능별 실무반 · 유관기관)을 따른다.", 64),
    ("구현 상태 범례", "구현 = 프로토타입에서 동작 확인  |  부분 = 구현했으나 외부 연계 실검증 미완(UNI 사내망)  |  모의 = 실 API 대신 Seed/Mock 으로 동일 인터페이스 제공(T3Q·기상·SMS 실발송)  |  계획 = 자리만 확보(KRMS 등)", 48),
    ("변경 표시", f"연두색(요구사항 ID) = v0.7 신규 {n_new}건  |  연노랑 = v0.7 수정·보완 {n_mod}건  |  총 {len(ROWS)}건 (구현 {st.get('구현',0)} · 부분 {st.get('부분',0)} · 모의 {st.get('모의',0)} · 계획 {st.get('계획',0)})", 32),
    ("v0.7 주요 변경", "① SOP 라이브러리(상황 독립 작성·게시·배포·즉시 실행)  ② 조직·연락처 3계층(재대본 편성)·엑셀 양식/일괄 업로드·전송그룹  ③ 상황전파 채널(SMS·이메일)·모바일 응답 페이지(수신확인→임무완료→조치사항)·실행내역/상황일지 자동 반영  ④ 세부행동 체크·조치결과 자동 정리·시작/종료 안내창·용어 도움말  ⑤ 시연모드(산불 15단계)  ⑥ 상황일지·결과보고 AI 초안을 UNI 챗으로, HWPX 행정문서 양식 내보내기  ⑦ 저장·배포 구조(DB 없음, Upstash 응답 저장소)", 96),
    ("미결·협의", "[미결협의] 시트 참조 — T3Q 실연계 규격, UNI 사내망 실검증, SMS·메일 실발송 연계, KRMS, 서버 DB 도입 여부, HWPX 한컴 실검증, 공개 저장소 보안", 32),
    ("문의", "㈜유엔이 (담당자:               / 연락처:               )", 22),
]
for i, (k, v, h) in enumerate(guide, start=4):
    a = ws.cell(i, 1, k); a.font = F_BOLD; a.fill = FILL_LABEL; a.border = BORDER; a.alignment = AL_TOP
    b = ws.cell(i, 2, v); b.font = F_BODY; b.border = BORDER; b.alignment = AL_TOP
    if k == "문의": b.fill = FILL_BEIGE
    ws.row_dimensions[i].height = h

# ── 시트 2: 요구사항 ──
ws2 = wb.create_sheet("요구사항")
H = ["번호", "기능영역", "요구사항 ID", "요구사항명", "요구사항 내용", "입력항목", "출력항목", "연계/처리", "구현 상태", "구현 화면·모듈", "비고 (변경이력)"]
W = [5, 13, 13, 26, 62, 24, 24, 16, 9, 40, 44]
header_row(ws2, 1, H, W)
ws2.freeze_panes = "A2"
for n, r in enumerate(ROWS, start=1):
    row = n + 1
    vals = [n, r[1], r[2], r[3], r[4], r[5], r[6], r[7], r[8], r[9], r[10]]
    for c, v in enumerate(vals, start=1):
        cell = ws2.cell(row, c, v); cell.border = BORDER; cell.font = F_BODY; cell.alignment = AL_TOP
    ws2.cell(row, 1).font = F_GRAY; ws2.cell(row, 1).alignment = AL_RIGHT
    ws2.cell(row, 3).alignment = AL_TOP_NOWRAP
    ws2.cell(row, 4).font = F_BOLD
    ws2.cell(row, 9).fill = FILL_BEIGE; ws2.cell(row, 9).alignment = AL_CENTER
    if r[11] == "new": ws2.cell(row, 3).fill = FILL_NEW
    elif r[11] == "mod": ws2.cell(row, 3).fill = FILL_MOD
    ws2.row_dimensions[row].height = est_height([r[3], r[4], r[5], r[6], r[9], r[10]], [W[3], W[4], W[5], W[6], W[9], W[10]])
ws2.auto_filter.ref = f"A1:{get_column_letter(len(H))}{len(ROWS)+1}"

# ── 시트 3: 미결협의 ──
ws3 = wb.create_sheet("미결협의")
H3 = ["번호", "구분", "확인·협의 요청 내용", "필요 시점", "관련 요구사항 ID", "회신", "의견"]
W3 = [5, 16, 70, 14, 26, 9, 34]
header_row(ws3, 1, H3, W3); ws3.freeze_panes = "A2"
ISSUES = [
    ("T3Q", "문서조회·SOP 추천·기상요약 실 API 규격(엔드포인트·응답 스키마) 확정 및 Mock 어댑터(lib/t3q/adapter) 교체 일정. 현재는 동일 인터페이스의 Seed 기반 Mock.", "연계 협의 시", "UFR-002-002, UFR-003-001, UFR-004-001"),
    ("UNI RAG", "사내망(10.20.10.101) 에서 /chat/json(SOP)·일반 챗(/chat 가정) 실호출 검증. Vercel 등 외부에서 사용하려면 공인 접근 경로 또는 프록시 필요.", "즉시", "UFR-004-012, UFR-007-003, UFR-008-006"),
    ("SMS·이메일", "UNE 기존 SMS 모듈·메일 서버 실발송 연계(발신번호·문안 90자 제한·URL 단축). 현재는 모의 발송 + 링크·QR·문안 생성.", "2차 개발", "UFR-006-002"),
    ("KRMS", "재난관리자원통합관리시스템 API 연계 범위·일정. 현재는 UI 자리(비활성 탭)만 확보.", "향후", "UFR-006-005"),
    ("저장구조", "여러 상황실 PC 가 같은 상황을 공유·동시 편집해야 하는지 여부 → 필요 시 서버 DB(Postgres 등) 도입. 현재는 브라우저 localStorage + JSON 백업.", "요구 확인", "UFR-001-005, UFR-009-007"),
    ("현장 응답 저장소", "Upstash Redis(Vercel 마켓플레이스) 운영 계정·보관기간(현재 7일)·토큰 재발급 정책. 시연 후 토큰 재발급 권장.", "운영 전", "UFR-006-011"),
    ("문서 양식", "생성 HWPX 의 한컴오피스 실제 열기·인쇄 확인. 지자체 결재 양식(templete) 추가 반영 여부.", "시연 전", "UFR-008-008"),
    ("조직 데이터", "지자체 실 재대본 편성표·연락망(엑셀) 확보 → 일괄 업로드로 Seed 교체. 조치별 주관/지원/협업 부서명과 조직 부서명 표기 일치 여부 확인.", "시연 전", "UFR-006-007, UFR-006-008"),
    ("모바일", "iPhone·Safari 실기기 확인(표준 CSS 만 사용, 미검증). 갤럭시·크롬은 뷰포트 390px 자동 검증 완료.", "시연 전", "UFR-006-010"),
    ("보안", "GitHub 저장소가 공개 상태. 비공개 전환 및 .env 관리 정책 확인. 모바일 링크는 인증 없음(사용자 전제) → 링크 외부 유출 주의 안내 포함.", "즉시", "UFR-009-007, UFR-006-010"),
    ("문서 반영", "요구정의 v0.6/시나리오 v1.0/흐름도 v1.1 갱신본(2026-09-13) 과 개발 결과 대조 → 본 v0.7/v1.1/v1.2 로 반영. 잔여 차이는 검토 의견란에 기재 요청.", "검토 시", "-"),
]
for n, (k, c, t, ids) in enumerate(ISSUES, start=1):
    row = n + 1
    vals = [n, k, c, t, ids, None, None]
    for ci, v in enumerate(vals, start=1):
        cell = ws3.cell(row, ci, v); cell.border = BORDER; cell.font = F_BODY; cell.alignment = AL_TOP
    ws3.cell(row, 1).font = F_GRAY; ws3.cell(row, 1).alignment = AL_RIGHT
    ws3.cell(row, 2).font = F_BOLD
    ws3.cell(row, 4).alignment = AL_CENTER
    ws3.cell(row, 5).alignment = AL_TOP_NOWRAP
    ws3.cell(row, 6).fill = FILL_BEIGE; ws3.cell(row, 6).alignment = AL_CENTER
    ws3.row_dimensions[row].height = est_height([c], [W3[2]])

wb.save(OUT)
print("saved", OUT)
